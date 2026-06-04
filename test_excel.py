import openpyxl
import pandas as pd
import numpy as np
file_path = "./test.xlsx"
def write_data_by_row_col_name(wb, sheet_name, row_name, column_name, value):
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Worksheet '{sheet_name}' does not exist")
    ws = wb[sheet_name]

    target_col_idx = None
    for col in range(1, ws.max_column + 1):
        cell_value = ws.cell(row=1, column=col).value
        if cell_value == column_name:
            target_col_idx = col
            break
    if target_col_idx is None:
        raise ValueError(f"Column name '{column_name}' not found in the header (first row)")

    target_row_idx = None
    for row in range(1, ws.max_row + 1):
        cell_value = ws.cell(row=row, column=1).value
        if cell_value == row_name:
            target_row_idx = row
            break
    if target_row_idx is None:
        raise ValueError(f"Row name '{row_name}' not found in the first column")

    ws.cell(row=target_row_idx, column=target_col_idx, value=value)


def read_excel_col_to_array(sheet_name: str | int, col_name: str) -> np.ndarray:
    try:
        print(f"Reading file: {file_path}, Sheet: {sheet_name}...")
        df = pd.read_excel(file_path, sheet_name=sheet_name)

        if col_name not in df.columns:
            raise ValueError(f"Column name '{col_name}' is not found in the sheet. Available columns: {list(df.columns)}")

        col_array = df[col_name].to_numpy()

        print(f"✅ Successfully extracted column '{col_name}', array shape: {col_array.shape}, data type: {col_array.dtype}")
        return col_array

    except FileNotFoundError:
        raise FileNotFoundError(f"File not found: {file_path}")
    except ValueError as ve:
        raise ve
    except Exception as e:
        raise RuntimeError(f"Failed to read: {e}")